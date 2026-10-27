from docx import Document

doc = Document('Academic_Report.docx')
print('=== DOCX VERIFICATION ===')
print('Total paragraphs:', len(doc.paragraphs))
print('Total tables:', len(doc.tables))
table_shapes = [(len(t.rows), len(t.columns)) for t in doc.tables]
print('Table shapes:', table_shapes[:5])
print()

h_counts = {}
for p in doc.paragraphs:
    if p.style.name.startswith('Heading'):
        h_counts[p.style.name] = h_counts.get(p.style.name, 0) + 1
print('Headings:', h_counts)
print()

key_sections = ['1. Abstract', '2. Introduction', '3. Requirements', '4. System Design',
                '5. Data Flow', '6. Functionalities', '7. Implementation Details',
                '8. Source Code', '9. Proof and Screenshots', '10. Results',
                '11. Conclusion', 'Appendix']
full_text = '\n'.join(p.text for p in doc.paragraphs)
for section in key_sections:
    found = section in full_text
    status = 'OK' if found else 'MISSING'
    print('  ' + status + '  ' + section)

print()
print('File size:', len(open('Academic_Report.docx', 'rb').read()), 'bytes')
