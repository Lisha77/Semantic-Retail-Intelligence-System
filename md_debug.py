# Debug version of md_to_docx with progress printing
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HEX_NAVY = '1F3B63'
HEX_BLUE = '2E558A'
HEX_LIGHT = 'F0F0F0'

def strip_markup(text):
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'__(.*?)__', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'`([^`]*)`', r'\1', text)
    return text

def style_code(p):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), HEX_LIGHT)
    pPr.append(shd)
    for run in p.runs:
        run.font.name = 'Consolas'
        run.font.size = Pt(9)

def setup_styles(doc):
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    for lvl, size, color in [(1, 16, HEX_NAVY), (2, 14, HEX_BLUE), (3, 12, '3A6BA5')]:
        h = doc.styles['Heading %d' % lvl]
        h.font.name = 'Calibri'
        h.font.size = Pt(size)
        h.font.color.rgb = RGBColor.from_string(color)

def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '888888')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def is_table_sep(line):
    s = line.strip()
    if '|' not in s:
        return False
    inner = s.replace('|', '').replace('-', '').replace(' ', '').replace(':', '')
    return inner == ''

def convert(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.read().splitlines()

    doc = Document()
    setup_styles(doc)

    start = 0
    for i, ln in enumerate(lines):
        if ln.startswith('# SEMANTIC RETAIL INTELLIGENCE ENGINE'):
            start = i + 1
            break

    i = start
    n = len(lines)
    count = 0

    while i < n:
        count += 1
        if count % 50 == 0:
            print('Line', i, '/', n, flush=True)

        line = lines[i]
        s = line.strip()

        if i <= start + 1:
            i += 1
            continue

        if s.startswith('## '):
            text = s[3:]
            doc.add_heading(text, level=1)
            i += 1
            continue

        if s.startswith('### '):
            text = s[4:]
            doc.add_heading(text, level=2)
            i += 1
            continue

        if s == '---':
            add_hr(doc)
            i += 1
            continue

        if s == '```':
            code_lines = []
            i += 1
            fence_start = i
            while i < n and lines[i].strip() != '```':
                code_lines.append(lines[i])
                i += 1
            if i >= n:
                print('WARNING: unclosed code fence starting at line', fence_start)
            if code_lines:
                p = doc.add_paragraph()
                p.add_run('\n'.join(code_lines))
                style_code(p)
                doc.add_paragraph()
            i += 1
            continue

        if is_table_sep(s):
            print('TABLE at line', i, flush=True)
            table_lines = [s]
            i += 1
            inner_count = 0
            while i < n:
                inner_count += 1
                if inner_count > 100:
                    print('WARNING: inner loop too long at line', i)
                    break
                cur = lines[i]
                cs = cur.strip()
                if '|' not in cs or is_table_sep(cs) or cur.startswith('## ') or cur.startswith('### ') or cs == '```' or cs == '':
                    break
                table_lines.append(lines[i])
                i += 1
            table_lines = [t for t in table_lines if not is_table_sep(t)]
            if len(table_lines) >= 2:
                rows = []
                for tl in table_lines:
                    rows.append([c.strip() for c in tl.split('|')])
                rows = [r for r in rows if any(c for c in r)]
                if len(rows) >= 2:
                    ncols = len(rows[0])
                    table = doc.add_table(rows=len(rows), cols=ncols)
                    table.style = 'Light Grid Accent 1'
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for r_idx, row in enumerate(rows):
                        for c_idx, cell_text in enumerate(row):
                            cell = table.cell(r_idx, c_idx)
                            cell.text = cell_text
                            for p in cell.paragraphs:
                                p.paragraph_format.space_after = Pt(2)
                                p.paragraph_format.space_before = Pt(2)
                                for run in p.runs:
                                    run.font.size = Pt(9)
                                    run.font.name = 'Calibri'
                    doc.add_paragraph()
            continue

        if re.match(r'^[-*+]\s+|\d+\.?\s+', s):
            items = []
            while i < n and re.match(r'^[-*+]\s+|\d+\.?\s+', lines[i]):
                items.append(lines[i].lstrip(' -*+1234567890.'))
                i += 1
            for item in items:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(strip_markup(item.strip()))
            continue

        if s == '':
            i += 1
            continue
        p = doc.add_paragraph()
        p.add_run(strip_markup(line))
        i += 1

    doc.save(docx_path)
    print('Saved:', docx_path)

if __name__ == '__main__':
    convert('Academic_Report.md', 'Academic_Report.docx')
    print('Done.')
